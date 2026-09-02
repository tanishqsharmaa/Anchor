"""
test_pptx_renderer.py — Unit Tests for Deterministic PPTX Renderer
"""

from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches

from anchor.deck.ast_compiler import ASTCompiler
from anchor.deck.pptx_renderer import PPTXRenderer, SLIDE_WIDTH, SLIDE_HEIGHT
from anchor.deck.schemas import (
    SlideType,
    HeroSlideContent,
    BLUFSlideContent,
    PolicyMatrixSlideContent,
    MatrixRow,
    FinancialDelegationSlideContent,
    FinancialTierItem,
    SlideItem,
    SlideDeckAST,
)


def test_pptx_renderer_output_dimensions_and_slides(tmp_path):
    renderer = PPTXRenderer(output_dir=tmp_path)

    # Build 4-slide AST covering all archetypes
    hero = SlideItem(
        slide_id="s1",
        type=SlideType.HERO_SLIDE,
        title="Cover",
        content=HeroSlideContent(
            title="TACTICAL DRONE PROCUREMENT",
            subtitle="Western Fleet Replenishment",
            dtg="310300Z AUG 2026",
            officer="CDR A. SHARMA",
            unit="HQ WNC",
        ),
    )
    bluf = SlideItem(
        slide_id="s2",
        type=SlideType.BLUF_EXECUTIVE,
        title="BLUF Summary",
        content=BLUFSlideContent(
            bluf_headline="Sanction of ₹18.00 Cr required under DFPDS Schedule 7.",
            key_takeaways=["Takeaway 1", "Takeaway 2", "Takeaway 3"],
            decision_requested="Fleet Commander financial sanction.",
            risk_summary="Supply chain lead time 45 days.",
        ),
    )
    matrix = SlideItem(
        slide_id="s3",
        type=SlideType.POLICY_MATRIX,
        title="Policy Matrix",
        content=PolicyMatrixSlideContent(
            matrix_title="Mode Comparison",
            headers=["Parameter", "OTE", "STE"],
            rows=[
                MatrixRow(row_title="Bids", cells=["Min 3", "PAC only"]),
                MatrixRow(row_title="IFA", cells=["> Tier 4", "Mandatory"]),
            ],
            statutory_precedence_note="DFPDS-2026 prevails.",
        ),
    )
    fin = SlideItem(
        slide_id="s4",
        type=SlideType.FINANCIAL_DELEGATION,
        title="Financial Limits",
        content=FinancialDelegationSlideContent(
            schedule_title="Schedule 7: Tactical Drones",
            schedule_no=7,
            tiers=[
                FinancialTierItem(
                    tier="Tier 1",
                    tier_name="CNS",
                    with_ifa_limit=50.0,
                    without_ifa_limit=5.0,
                ),
                FinancialTierItem(
                    tier="Tier 3",
                    tier_name="Fleet Commander",
                    with_ifa_limit=18.0,
                    without_ifa_limit=1.0,
                ),
            ],
            notes="Mandatory IFA concurrence applies.",
        ),
    )

    deck_ast = SlideDeckAST(
        deck_id="deck_test_pptx_01",
        title="Test Briefing Presentation",
        topic="Tactical Drones",
        dtg="310300Z AUG 2026",
        slides=[hero, bluf, matrix, fin],
    )

    out_file = renderer.render(deck_ast)
    assert out_file.exists()
    assert out_file.suffix == ".pptx"

    # Re-open with python-pptx to verify validity and geometry
    prs = Presentation(str(out_file))
    assert len(prs.slides) == 4
    assert prs.slide_width == SLIDE_WIDTH
    assert prs.slide_height == SLIDE_HEIGHT


def test_pptx_renderer_deterministic_schedule_deck(tmp_path):
    compiler = ASTCompiler()
    deck_ast = compiler.compile_deterministic_schedule(schedule_no=18)

    renderer = PPTXRenderer(output_dir=tmp_path)
    out_file = renderer.render(deck_ast, output_filename="schedule_18_briefing.pptx")

    assert out_file.exists()
    prs = Presentation(str(out_file))
    assert len(prs.slides) == deck_ast.total_slides
