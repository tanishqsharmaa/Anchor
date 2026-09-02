"""
pptx_renderer.py — Deterministic PowerPoint Presentation Renderer for AutoDeck AI
"""

from pathlib import Path
from typing import Optional, Union
import logging

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from anchor.config import settings
from anchor.deck.schemas import (
    SlideType,
    SlideDeckAST,
    SlideItem,
    HeroSlideContent,
    BLUFSlideContent,
    PolicyMatrixSlideContent,
    FinancialDelegationSlideContent,
)

logger = logging.getLogger(__name__)

# Naval Briefing 16:9 Widescreen Geometry
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Naval Presentation Color Palette
COLOR_BG_NAVY = RGBColor(10, 25, 47)       # #0A192F Deep Naval Blue
COLOR_CARD_SURFACE = RGBColor(17, 34, 64)  # #112240 Dark Surface Card
COLOR_TEXT_PRIMARY = RGBColor(255, 255, 255) # #FFFFFF White
COLOR_TEXT_BODY = RGBColor(204, 214, 246)  # #CCD6F6 Ice Blue
COLOR_TEXT_MUTED = RGBColor(136, 146, 176) # #8892B0 Muted Slate
COLOR_GOLD = RGBColor(255, 215, 0)         # #FFD700 Naval Gold Accent
COLOR_CYAN = RGBColor(0, 229, 255)         # #00E5FF Tactical Cyan
COLOR_BANNER_RED = RGBColor(239, 68, 68)   # #EF4444 Classification Red
COLOR_BORDER = RGBColor(35, 53, 84)        # #233554 Card Border

# Typography Standards
FONT_TITLE = "Arial"
FONT_BODY = "Segoe UI"
FONT_MONO = "Consolas"


class PPTXRenderer:
    """
    Deterministic PowerPoint presentation generator compiling SlideDeckAST into 16:9 widescreen slides.
    Adheres strictly to Indian Naval briefing typography, classification marking, and character budgets.
    """

    def __init__(self, output_dir: Optional[Union[Path, str]] = None) -> None:
        if output_dir:
            self.output_dir = Path(output_dir)
        else:
            self.output_dir = settings.DATA_DIR / "decks"
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def render(self, deck_ast: SlideDeckAST, output_filename: Optional[str] = None) -> Path:
        """
        Compile SlideDeckAST into a binary .pptx file.

        Args:
            deck_ast: Validated SlideDeckAST model.
            output_filename: Optional target filename (default {deck_id}.pptx).

        Returns:
            Path: Absolute path to the generated .pptx file.
        """
        prs = Presentation()
        prs.slide_width = SLIDE_WIDTH
        prs.slide_height = SLIDE_HEIGHT

        blank_layout = prs.slide_layouts[6]  # Completely blank slide layout

        for slide_idx, slide_item in enumerate(deck_ast.slides):
            slide = prs.slides.add_slide(blank_layout)
            self._apply_background(slide)
            self._apply_classification_banners(slide, slide_item.classification or deck_ast.classification)

            if slide_item.type == SlideType.HERO_SLIDE:
                self._render_hero_slide(slide, slide_item)
            elif slide_item.type == SlideType.BLUF_EXECUTIVE:
                self._render_bluf_slide(slide, slide_item)
            elif slide_item.type == SlideType.POLICY_MATRIX:
                self._render_policy_matrix_slide(slide, slide_item)
            elif slide_item.type == SlideType.FINANCIAL_DELEGATION:
                self._render_financial_delegation_slide(slide, slide_item)
            else:
                self._render_generic_slide(slide, slide_item)

            self._apply_footer(slide, slide_idx + 1, len(deck_ast.slides), deck_ast.dtg)

        fname = output_filename or f"{deck_ast.deck_id}.pptx"
        if not fname.endswith(".pptx"):
            fname += ".pptx"

        out_path = self.output_dir / fname
        prs.save(str(out_path))
        logger.info(f"[PPTXRenderer] Successfully rendered deck to {out_path}")
        return out_path

    def _apply_background(self, slide) -> None:
        """Add dark naval background rectangle covering entire slide."""
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            SLIDE_WIDTH,
            SLIDE_HEIGHT,
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG_NAVY
        bg_shape.line.fill.background()

    def _apply_classification_banners(self, slide, classification_text: str) -> None:
        """Add standard top and bottom security classification banners."""
        # Top Banner
        top_box = slide.shapes.add_textbox(Inches(0), Inches(0.1), SLIDE_WIDTH, Inches(0.3))
        tf_top = top_box.text_frame
        tf_top.word_wrap = True
        p_top = tf_top.paragraphs[0]
        p_top.text = classification_text.upper()
        p_top.alignment = PP_ALIGN.CENTER
        p_top.font.name = FONT_TITLE
        p_top.font.size = Pt(9.5)
        p_top.font.bold = True
        p_top.font.color.rgb = COLOR_BANNER_RED

        # Bottom Banner
        bot_box = slide.shapes.add_textbox(Inches(0), Inches(7.1), SLIDE_WIDTH, Inches(0.3))
        tf_bot = bot_box.text_frame
        tf_bot.word_wrap = True
        p_bot = tf_bot.paragraphs[0]
        p_bot.text = classification_text.upper()
        p_bot.alignment = PP_ALIGN.CENTER
        p_bot.font.name = FONT_TITLE
        p_bot.font.size = Pt(9.5)
        p_bot.font.bold = True
        p_bot.font.color.rgb = COLOR_BANNER_RED

    def _apply_footer(self, slide, page_num: int, total_pages: int, dtg: Optional[str]) -> None:
        """Add slide number and DTG footer readouts."""
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.3))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.font.name = FONT_MONO
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_TEXT_MUTED

        footer_text = f"PROJECT ANCHOR // PAGE {page_num:02d} OF {total_pages:02d}"
        if dtg:
            footer_text = f"{dtg}  |  {footer_text}"
        p.text = footer_text

    def _render_hero_slide(self, slide, item: SlideItem) -> None:
        """Render Title / Hero presentation slide."""
        content: HeroSlideContent = (
            item.content if isinstance(item.content, HeroSlideContent)
            else HeroSlideContent.model_validate(item.content)
        )

        # Card container
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1.2),
            Inches(1.2),
            Inches(10.933),
            Inches(5.0),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_SURFACE
        card.line.color.rgb = COLOR_CYAN
        card.line.width = Pt(1.5)

        # Header Badge
        badge_box = slide.shapes.add_textbox(Inches(1.6), Inches(1.6), Inches(10.133), Inches(0.4))
        p_badge = badge_box.text_frame.paragraphs[0]
        p_badge.text = "INDIAN NAVAL STAFF BRIEFING  |  OPERATIONAL COMPLIANCE"
        p_badge.font.name = FONT_MONO
        p_badge.font.size = Pt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = COLOR_CYAN

        # Title Box
        title_box = slide.shapes.add_textbox(Inches(1.6), Inches(2.2), Inches(10.133), Inches(1.4))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = content.title.upper()
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_PRIMARY

        # Subtitle Box
        if content.subtitle:
            p_sub = tf_title.add_paragraph()
            p_sub.text = content.subtitle
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(18)
            p_sub.font.color.rgb = COLOR_GOLD
            p_sub.space_before = Pt(8)

        # Metadata Details (Officer / Unit / DTG)
        meta_box = slide.shapes.add_textbox(Inches(1.6), Inches(4.3), Inches(10.133), Inches(1.4))
        tf_meta = meta_box.text_frame
        tf_meta.word_wrap = True

        p_off = tf_meta.paragraphs[0]
        p_off.text = f"PRESENTER:  {content.officer or 'STAFF OFFICER (OPERATIONS & COMPLIANCE)'}"
        p_off.font.name = FONT_BODY
        p_off.font.size = Pt(13)
        p_off.font.bold = True
        p_off.font.color.rgb = COLOR_TEXT_PRIMARY

        p_unit = tf_meta.add_paragraph()
        p_unit.text = f"COMMAND:    {content.unit or 'INTEGRATED HEADQUARTERS, MINISTRY OF DEFENCE (NAVY)'}"
        p_unit.font.name = FONT_BODY
        p_unit.font.size = Pt(12)
        p_unit.font.color.rgb = COLOR_TEXT_BODY

        if content.dtg:
            p_dtg = tf_meta.add_paragraph()
            p_dtg.text = f"DTG:        {content.dtg}"
            p_dtg.font.name = FONT_MONO
            p_dtg.font.size = Pt(11)
            p_dtg.font.color.rgb = COLOR_TEXT_MUTED

    def _render_bluf_slide(self, slide, item: SlideItem) -> None:
        """Render Bottom Line Up Front (BLUF) executive decision slide."""
        content: BLUFSlideContent = (
            item.content if isinstance(item.content, BLUFSlideContent)
            else BLUFSlideContent.model_validate(item.content)
        )

        # Slide Heading
        self._add_slide_header(slide, item.title or "EXECUTIVE SUMMARY (BLUF)")

        # Left Card: Headline & Takeaways
        left_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8),
            Inches(1.4),
            Inches(6.8),
            Inches(5.0),
        )
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = COLOR_CARD_SURFACE
        left_card.line.color.rgb = COLOR_BORDER

        # Left Card Content
        tf_left = left_card.text_frame
        tf_left.word_wrap = True
        tf_left.vertical_anchor = MSO_ANCHOR.TOP
        tf_left.margin_left = Inches(0.3)
        tf_left.margin_top = Inches(0.3)
        tf_left.margin_right = Inches(0.3)

        p_bluf_lbl = tf_left.paragraphs[0]
        p_bluf_lbl.text = "BOTTOM LINE UP FRONT"
        p_bluf_lbl.font.name = FONT_MONO
        p_bluf_lbl.font.size = Pt(11)
        p_bluf_lbl.font.bold = True
        p_bluf_lbl.font.color.rgb = COLOR_GOLD

        p_head = tf_left.add_paragraph()
        p_head.text = content.bluf_headline
        p_head.font.name = FONT_TITLE
        p_head.font.size = Pt(16)
        p_head.font.bold = True
        p_head.font.color.rgb = COLOR_TEXT_PRIMARY
        p_head.space_before = Pt(6)
        p_head.space_after = Pt(12)

        p_take_lbl = tf_left.add_paragraph()
        p_take_lbl.text = "KEY TAKEAWAYS & FINDINGS:"
        p_take_lbl.font.name = FONT_MONO
        p_take_lbl.font.size = Pt(10)
        p_take_lbl.font.bold = True
        p_take_lbl.font.color.rgb = COLOR_CYAN
        p_take_lbl.space_after = Pt(6)

        for takeaway in content.key_takeaways:
            p_item = tf_left.add_paragraph()
            p_item.text = f"• {takeaway}"
            p_item.font.name = FONT_BODY
            p_item.font.size = Pt(13)
            p_item.font.color.rgb = COLOR_TEXT_BODY
            p_item.space_before = Pt(4)

        # Right Card: Decision Requested & Risk Summary
        right_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(7.8),
            Inches(1.4),
            Inches(4.733),
            Inches(5.0),
        )
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = COLOR_CARD_SURFACE
        right_card.line.color.rgb = COLOR_BORDER

        tf_right = right_card.text_frame
        tf_right.word_wrap = True
        tf_right.vertical_anchor = MSO_ANCHOR.TOP
        tf_right.margin_left = Inches(0.3)
        tf_right.margin_top = Inches(0.3)
        tf_right.margin_right = Inches(0.3)

        p_urg = tf_right.paragraphs[0]
        p_urg.text = f"URGENCY: {content.urgency.value}"
        p_urg.font.name = FONT_MONO
        p_urg.font.size = Pt(11)
        p_urg.font.bold = True
        p_urg.font.color.rgb = COLOR_CYAN

        p_dec_lbl = tf_right.add_paragraph()
        p_dec_lbl.text = "DECISION / SANCTION REQUESTED:"
        p_dec_lbl.font.name = FONT_TITLE
        p_dec_lbl.font.size = Pt(13)
        p_dec_lbl.font.bold = True
        p_dec_lbl.font.color.rgb = COLOR_GOLD
        p_dec_lbl.space_before = Pt(10)

        p_dec = tf_right.add_paragraph()
        p_dec.text = content.decision_requested
        p_dec.font.name = FONT_BODY
        p_dec.font.size = Pt(13)
        p_dec.font.color.rgb = COLOR_TEXT_PRIMARY
        p_dec.space_before = Pt(4)

        if content.risk_summary:
            p_risk_lbl = tf_right.add_paragraph()
            p_risk_lbl.text = "OPERATIONAL & AUDIT RISK:"
            p_risk_lbl.font.name = FONT_TITLE
            p_risk_lbl.font.size = Pt(13)
            p_risk_lbl.font.bold = True
            p_risk_lbl.font.color.rgb = COLOR_BANNER_RED
            p_risk_lbl.space_before = Pt(16)

            p_risk = tf_right.add_paragraph()
            p_risk.text = content.risk_summary
            p_risk.font.name = FONT_BODY
            p_risk.font.size = Pt(12.5)
            p_risk.font.color.rgb = COLOR_TEXT_BODY
            p_risk.space_before = Pt(4)

    def _render_policy_matrix_slide(self, slide, item: SlideItem) -> None:
        """Render Policy Matrix comparative table slide."""
        content: PolicyMatrixSlideContent = (
            item.content if isinstance(item.content, PolicyMatrixSlideContent)
            else PolicyMatrixSlideContent.model_validate(item.content)
        )

        self._add_slide_header(slide, item.title or "REGULATORY POLICY MATRIX")

        # Table dimensions
        num_rows = len(content.rows) + 1
        num_cols = len(content.headers)

        table_shape = slide.shapes.add_table(
            num_rows,
            num_cols,
            Inches(0.8),
            Inches(1.5),
            Inches(11.733),
            Inches(0.6 * num_rows),
        )
        table = table_shape.table

        # Headers
        for c_idx, h_text in enumerate(content.headers):
            cell = table.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_SURFACE
            cell.text_frame.word_wrap = True
            p = cell.text_frame.paragraphs[0]
            p.text = h_text.upper()
            p.font.name = FONT_TITLE
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = COLOR_CYAN
            p.alignment = PP_ALIGN.LEFT

        # Data Rows
        for r_idx, row in enumerate(content.rows):
            # If row.cells matches num_cols - 1, cell 0 is row_title
            if len(row.cells) == num_cols - 1:
                row_items = [row.row_title] + row.cells
            else:
                row_items = row.cells

            for c_idx, val in enumerate(row_items[:num_cols]):
                cell = table.cell(r_idx + 1, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_BG_NAVY
                cell.text_frame.word_wrap = True
                p = cell.text_frame.paragraphs[0]
                p.text = val
                p.font.name = FONT_BODY
                p.font.size = Pt(12)
                p.font.color.rgb = COLOR_TEXT_PRIMARY if c_idx == 0 else COLOR_TEXT_BODY
                if c_idx == 0:
                    p.font.bold = True

        # Precedence Note Box
        if content.statutory_precedence_note:
            note_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.5))
            tf = note_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"STATUTORY PRECEDENCE: {content.statutory_precedence_note}"
            p.font.name = FONT_MONO
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_GOLD

    def _render_financial_delegation_slide(self, slide, item: SlideItem) -> None:
        """Render Financial Delegation Schedule breakdown slide."""
        content: FinancialDelegationSlideContent = (
            item.content if isinstance(item.content, FinancialDelegationSlideContent)
            else FinancialDelegationSlideContent.model_validate(item.content)
        )

        self._add_slide_header(slide, item.title or f"DFPDS SCHEDULE {content.schedule_no:02d} DELEGATIONS")

        # Table
        headers = ["CFA Tier", "Competent Financial Authority", "With IFA Concurrence", "Without IFA", "PAC Limit"]
        num_rows = len(content.tiers) + 1
        num_cols = len(headers)

        table_shape = slide.shapes.add_table(
            num_rows,
            num_cols,
            Inches(0.8),
            Inches(1.5),
            Inches(11.733),
            Inches(0.55 * num_rows),
        )
        table = table_shape.table

        # Column Widths
        table.columns[0].width = Inches(1.8)
        table.columns[1].width = Inches(3.933)
        table.columns[2].width = Inches(2.2)
        table.columns[3].width = Inches(1.8)
        table.columns[4].width = Inches(2.0)

        # Header Row
        for c_idx, h in enumerate(headers):
            cell = table.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_SURFACE
            cell.text_frame.word_wrap = True
            p = cell.text_frame.paragraphs[0]
            p.text = h.upper()
            p.font.name = FONT_TITLE
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_GOLD
            if c_idx >= 2:
                p.alignment = PP_ALIGN.RIGHT

        # Tiers
        for r_idx, tier in enumerate(content.tiers):
            row_idx = r_idx + 1
            row_vals = [
                tier.tier,
                tier.tier_name,
                f"₹{tier.with_ifa_limit:.2f} Cr" if tier.with_ifa_limit > 0 else "Nil / Full",
                f"₹{tier.without_ifa_limit:.2f} Cr" if tier.without_ifa_limit > 0 else "Nil",
                f"₹{tier.pac_limit:.2f} Cr" if tier.pac_limit is not None and tier.pac_limit > 0 else "—",
            ]

            for c_idx, val in enumerate(row_vals):
                cell = table.cell(row_idx, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_BG_NAVY
                cell.text_frame.word_wrap = True
                p = cell.text_frame.paragraphs[0]
                p.text = val
                p.font.name = FONT_BODY
                p.font.size = Pt(12)
                p.font.color.rgb = COLOR_TEXT_PRIMARY if c_idx < 2 else COLOR_CYAN
                if c_idx >= 2:
                    p.alignment = PP_ALIGN.RIGHT
                    p.font.bold = True

        # Bottom Notes
        if content.notes:
            notes_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.5))
            tf = notes_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"STATUTORY PROVISO: {content.notes}"
            p.font.name = FONT_MONO
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_TEXT_MUTED

    def _render_generic_slide(self, slide, item: SlideItem) -> None:
        """Render generic fallback slide for custom contents."""
        self._add_slide_header(slide, item.title)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8),
            Inches(1.4),
            Inches(11.733),
            Inches(5.0),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_SURFACE
        card.line.color.rgb = COLOR_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_top = Inches(0.4)
        p = tf.paragraphs[0]
        p.text = str(item.content)
        p.font.name = FONT_BODY
        p.font.size = Pt(14)
        p.font.color.rgb = COLOR_TEXT_BODY

    def _add_slide_header(self, slide, title_text: str) -> None:
        """Helper to add standard slide title header."""
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.7))
        tf = header_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text.upper()
        p.font.name = FONT_TITLE
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY
